from __future__ import annotations

import io
import tempfile
import unittest
from contextlib import redirect_stderr
from pathlib import Path

from docx import Document as load_docx_document

from resume_builder.models import (
    ExperienceEntry,
    ExperienceType,
    ResumeContent,
    ResumeMeta,
    SectionKind,
    SectionTitles,
)
from resume_builder.parser import parse_resume_source
from resume_builder.profiles import OutputProfile
from resume_builder.renderer import (
    build_doc_from_source,
    build_markdown_from_source,
    render_experience_markdown,
    render_professional_experience_section,
)
from resume_builder.sections import SECTION_DEFINITIONS, load_resume_directory
from tests.helpers import (
    RESUME_SOURCE,
    copy_resume_source,
    heading_texts,
    read_frontmatter_title,
    read_section_titles,
)


class ResumeRendererTests(unittest.TestCase):
    def test_canonical_source_builds_a_readable_docx_smoke_test(self) -> None:
        titles = read_section_titles(RESUME_SOURCE)

        with tempfile.TemporaryDirectory() as temp_dir:
            output = Path(temp_dir) / "resume.docx"
            stderr = io.StringIO()
            with redirect_stderr(stderr):
                document = build_doc_from_source(
                    RESUME_SOURCE, profile=OutputProfile.GROUPED
                )
                document.save(str(output))

            reloaded = load_docx_document(str(output))

        self.assertEqual(stderr.getvalue(), "")
        self.assertEqual(document.core_properties.title, "Davi Naizer Resume")
        self.assertEqual(document.core_properties.subject, "Resume")
        self.assertEqual(document.core_properties.author, "Davi Naizer")
        rendered_text = [paragraph.text for paragraph in reloaded.paragraphs]
        expected_experience_text = [
            "Gamesys → Bally's Interactive | London, UK\t03/2019",
            "Joined Gamesys as a Frontend Developer and progressed",
            "Frontend Tech Lead",
            "11/2022",
            "Provided technical direction across frontend initiatives",
            "Senior Frontend Engineer",
            "10/2020",
            "Worked across customer-facing products and internal engineering tooling",
            "Frontend Developer",
            "03/2019",
            "Joined the Promotions team building customer-facing",
            "Tech: React • TypeScript • Redux Toolkit",
            "Tech: React • JavaScript • Styled-components",
            "Planned Career Break",
            "Took a planned break from full-time employment",
        ]
        for text in expected_experience_text:
            self.assertTrue(any(text in paragraph for paragraph in rendered_text), text)
        role_headings = [
            "Frontend Tech Lead",
            "Senior Frontend Engineer",
            "Frontend Developer",
        ]
        role_indices = [
            next(
                index
                for index, paragraph in enumerate(rendered_text)
                if heading in paragraph
            )
            for heading in role_headings
        ]
        self.assertEqual(role_indices, sorted(role_indices))
        career_break_index = next(
            index
            for index, paragraph in enumerate(rendered_text)
            if paragraph.startswith("Planned Career Break\t")
        )
        next_entry_index = next(
            index
            for index, paragraph in enumerate(
                rendered_text[career_break_index + 1 :], career_break_index + 1
            )
            if paragraph.startswith(
                "Senior Frontend Software Engineer | The Signal Group"
            )
        )
        self.assertEqual(
            rendered_text[career_break_index + 1 : next_entry_index],
            [
                "Took a planned break from full-time employment after leaving The Signal Group. During that time I stepped back to re-evaluate my career direction and returned to hands-on product development, which eventually developed into the venture behind Alfred."
            ],
        )
        self.assertEqual(
            heading_texts(reloaded),
            [
                titles["summary"].upper(),
                titles["core_skills"].upper(),
                titles["professional_experience"].upper(),
                titles["education"].upper(),
            ],
        )

    def test_ats_profile_flattens_grouped_employment(self) -> None:
        document = build_doc_from_source(RESUME_SOURCE, profile=OutputProfile.ATS)
        paragraphs = document.paragraphs
        headings = [
            paragraph
            for paragraph in paragraphs
            if paragraph.style.name == "Heading 2"
            and any(
                title in paragraph.text
                for title in (
                    "Frontend Tech Lead",
                    "Senior Frontend Engineer",
                    "Frontend Developer",
                )
            )
        ]
        heading_texts = [paragraph.text for paragraph in headings]

        self.assertEqual(
            heading_texts,
            [
                "Frontend Tech Lead | Bally's Interactive | London, UK | 11/2022\N{NO-BREAK SPACE}–\N{NO-BREAK SPACE}11/2023",
                "Senior Frontend Engineer | Gamesys / Bally's Interactive | London, UK | 10/2020\N{NO-BREAK SPACE}–\N{NO-BREAK SPACE}11/2022",
                "Frontend Developer | Gamesys | London, UK | 03/2019\N{NO-BREAK SPACE}–\N{NO-BREAK SPACE}09/2020",
            ],
        )
        self.assertTrue(
            any(
                "Joined Gamesys as a Frontend Developer and progressed"
                in paragraph.text
                for paragraph in paragraphs
            )
        )
        self.assertTrue(
            all(paragraph.paragraph_format.left_indent == 0 for paragraph in headings)
        )
        self.assertFalse(
            any("Gamesys → Bally's Interactive" in text for text in heading_texts)
        )

    def test_ats_markdown_flattens_grouped_employment(self) -> None:
        markdown = build_markdown_from_source(RESUME_SOURCE, profile=OutputProfile.ATS)
        experience = markdown.split("# Professional Experience", 1)[1].split(
            "# Education", 1
        )[0]

        self.assertIn(
            "## **Frontend Tech Lead | Bally's Interactive | London, UK | 11/2022 – 11/2023**",
            experience,
        )
        self.assertIn(
            "## **Senior Frontend Engineer | Gamesys / Bally's Interactive | London, UK | 10/2020 – 11/2022**",
            experience,
        )
        self.assertIn(
            "## **Frontend Developer | Gamesys | London, UK | 03/2019 – 09/2020**",
            experience,
        )
        self.assertNotIn(
            "Gamesys → Bally's Interactive | London, UK | 03/2019 – 11/2023", experience
        )
        self.assertNotIn("### ", experience)
        self.assertNotIn("<!-- experience:", experience)
        self.assertIn(
            "Joined Gamesys as a Frontend Developer and progressed to Senior Frontend Engineer",
            experience,
        )

    def test_renderer_omits_invalid_career_break_bullets_and_technology(self) -> None:
        entry = ExperienceEntry(
            type=ExperienceType.CAREER_BREAK,
            title="Career Break",
            organisation=None,
            location=None,
            date_right="2025",
            descriptions=("Summary.",),
            bullets=("Must not render.",),
            tech="Must not render",
            roles=(),
        )
        content = ResumeContent(
            meta=ResumeMeta("Name", "Title", "Tagline", ("Contact",)),
            summary=(),
            skills=(),
            experience=(entry,),
            selected_project=None,
            education=(),
            section_titles=SectionTitles(),
            present_sections=frozenset({SectionKind.PROFESSIONAL_EXPERIENCE}),
            section_order=(SectionKind.PROFESSIONAL_EXPERIENCE,),
        )
        document = load_docx_document()

        render_professional_experience_section(document, content)

        self.assertEqual(
            [paragraph.text for paragraph in document.paragraphs],
            ["PROFESSIONAL EXPERIENCE", "Career Break\t2025", "Summary."],
        )
        for profile in (OutputProfile.ATS, OutputProfile.GROUPED):
            with self.subTest(profile=profile):
                markdown = render_experience_markdown((entry,), profile=profile)
                self.assertIn("Summary.", markdown)
                self.assertNotIn("Must not render.", markdown)

    def test_meta_section_order_controls_rendering_order(self) -> None:
        titles = read_section_titles(RESUME_SOURCE)
        with tempfile.TemporaryDirectory() as temp_dir:
            source = copy_resume_source(temp_dir)
            meta_path = source / "meta.md"
            meta_markdown = meta_path.read_text(encoding="utf-8").replace(
                "sections:\n  - summary\n  - core_skills\n  - professional_experience\n  - education\n",
                "sections:\n  - education\n  - summary\n  - core_skills\n  - professional_experience\n",
                1,
            )
            meta_path.write_text(meta_markdown, encoding="utf-8")

            content = parse_resume_source(source)
            document = build_doc_from_source(source)
            markdown = build_markdown_from_source(source)

        self.assertEqual(content.section_titles.education, titles["education"])
        self.assertEqual(
            heading_texts(document),
            [
                titles["education"].upper(),
                titles["summary"].upper(),
                titles["core_skills"].upper(),
                titles["professional_experience"].upper(),
            ],
        )
        self.assertLess(
            markdown.index(f"# {titles['education']}"),
            markdown.index(f"# {titles['summary']}"),
        )
        self.assertLess(
            markdown.index(f"# {titles['summary']}"),
            markdown.index(f"# {titles['core_skills']}"),
        )
        self.assertLess(
            markdown.index(f"# {titles['core_skills']}"),
            markdown.index(f"# {titles['professional_experience']}"),
        )

    def test_each_missing_required_section_is_not_rendered(self) -> None:
        required_definitions = [
            definition for definition in SECTION_DEFINITIONS if not definition.optional
        ]
        titles = read_section_titles(RESUME_SOURCE)
        expected_headings_by_filename = {
            "summary.md": titles["summary"].upper(),
            "core-skills.md": titles["core_skills"].upper(),
            "professional-experience.md": titles["professional_experience"].upper(),
            "education.md": titles["education"].upper(),
        }

        for missing_definition in required_definitions:
            with (
                self.subTest(filename=missing_definition.filename),
                tempfile.TemporaryDirectory() as temp_dir,
            ):
                source = copy_resume_source(temp_dir)
                missing_path = source / "sections" / missing_definition.filename
                missing_path.unlink()

                stderr = io.StringIO()
                with redirect_stderr(stderr):
                    document = build_doc_from_source(source)

                warning = stderr.getvalue()
                self.assertEqual(warning.count("Warning:"), 1)
                self.assertIn(str(missing_path), warning)
                expected_headings = list(expected_headings_by_filename.values())
                expected_headings.remove(
                    expected_headings_by_filename[missing_definition.filename]
                )
                self.assertEqual(heading_texts(document), expected_headings)

    def test_all_missing_required_sections_render_no_section_headings(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            source = copy_resume_source(temp_dir)
            required_definitions = [
                definition
                for definition in SECTION_DEFINITIONS
                if not definition.optional
            ]
            for definition in required_definitions:
                (source / "sections" / definition.filename).unlink()

            stderr = io.StringIO()
            with redirect_stderr(stderr):
                document = build_doc_from_source(source)

        warning = stderr.getvalue()
        self.assertEqual(warning.count("Warning:"), len(required_definitions))
        for definition in required_definitions:
            self.assertIn(str(source / "sections" / definition.filename), warning)
        self.assertEqual(heading_texts(document), [])

    def test_remaining_editable_title_is_preserved_after_an_earlier_omission(
        self,
    ) -> None:
        titles = read_section_titles(RESUME_SOURCE)
        with tempfile.TemporaryDirectory() as temp_dir:
            source = copy_resume_source(temp_dir)
            missing_path = source / "sections" / "core-skills.md"
            missing_path.unlink()
            education_path = source / "sections" / "education.md"
            education_title = read_frontmatter_title(education_path)
            education_markdown = education_path.read_text(encoding="utf-8").replace(
                f"title: {education_title}", "title: Learning & Credentials", 1
            )
            education_path.write_text(education_markdown, encoding="utf-8")

            parse_stderr = io.StringIO()
            with redirect_stderr(parse_stderr):
                content = parse_resume_source(source)
            render_stderr = io.StringIO()
            with redirect_stderr(render_stderr):
                document = build_doc_from_source(source)

        self.assertEqual(parse_stderr.getvalue().count("Warning:"), 1)
        self.assertEqual(render_stderr.getvalue().count("Warning:"), 1)
        self.assertIn(str(missing_path), parse_stderr.getvalue())
        self.assertIn(str(missing_path), render_stderr.getvalue())
        self.assertNotIn("core_skills", content.present_sections)
        self.assertEqual(content.skills, ())
        self.assertIn("summary", content.present_sections)
        self.assertIn("professional_experience", content.present_sections)
        self.assertIn("education", content.present_sections)
        self.assertEqual(content.section_titles.education, "Learning & Credentials")
        self.assertEqual(
            heading_texts(document),
            [
                titles["summary"].upper(),
                titles["professional_experience"].upper(),
                "LEARNING & CREDENTIALS",
            ],
        )

    def test_missing_optional_section_is_silent_and_not_rendered(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            source = copy_resume_source(temp_dir)

            stderr = io.StringIO()
            with redirect_stderr(stderr):
                content = parse_resume_source(source)
                document = build_doc_from_source(source)

        self.assertEqual(stderr.getvalue(), "")
        self.assertNotIn("selected_project", content.present_sections)
        self.assertIsNone(content.selected_project)
        self.assertNotIn("SELECTED PROJECT", heading_texts(document))

    def test_selected_project_is_rendered_with_its_editable_title(self) -> None:
        titles = read_section_titles(RESUME_SOURCE)
        markdown = """---
title: Selected Work
---

## Project | 2026

Description.

- Result.

**Tech:** Python
"""
        with tempfile.TemporaryDirectory() as temp_dir:
            source = copy_resume_source(temp_dir)
            meta_path = source / "meta.md"
            meta_markdown = meta_path.read_text(encoding="utf-8").replace(
                "sections:\n  - summary\n  - core_skills\n  - professional_experience\n  - education\n",
                "sections:\n  - summary\n  - core_skills\n  - professional_experience\n  - selected_project\n  - education\n",
                1,
            )
            meta_path.write_text(meta_markdown, encoding="utf-8")
            (source / "sections" / "selected-project.md").write_text(
                markdown, encoding="utf-8"
            )

            loaded_source = load_resume_directory(source)
            content = parse_resume_source(source)
            document = build_doc_from_source(source)

        selected_section = next(
            section
            for section in loaded_source.sections
            if section.kind == "selected_project"
        )
        self.assertEqual(selected_section.title, "Selected Work")
        self.assertIn("selected_project", content.present_sections)
        self.assertEqual(content.section_titles.selected_project, "Selected Work")
        self.assertIsNotNone(content.selected_project)
        assert content.selected_project is not None
        self.assertEqual(content.selected_project.heading_left, "Project")
        self.assertEqual(content.selected_project.bullets, ("Result.",))
        self.assertEqual(content.selected_project.tech, "Python")
        self.assertEqual(
            heading_texts(document),
            [
                titles["summary"].upper(),
                titles["core_skills"].upper(),
                titles["professional_experience"].upper(),
                "SELECTED WORK",
                titles["education"].upper(),
            ],
        )

    def test_build_markdown_combines_sections_in_order_with_editable_titles(
        self,
    ) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            source = copy_resume_source(temp_dir)
            meta_path = source / "meta.md"
            meta_markdown = meta_path.read_text(encoding="utf-8").replace(
                "sections:\n  - summary\n  - core_skills\n  - professional_experience\n  - education\n",
                "sections:\n  - summary\n  - core_skills\n  - professional_experience\n  - selected_project\n  - education\n",
                1,
            )
            meta_path.write_text(meta_markdown, encoding="utf-8")
            summary_path = source / "sections" / "summary.md"
            summary_title = read_frontmatter_title(summary_path)
            summary_markdown = summary_path.read_text(encoding="utf-8").replace(
                f"title: {summary_title}", "title: Professional Profile", 1
            )
            summary_path.write_text(summary_markdown, encoding="utf-8")
            selected_project_path = source / "sections" / "selected-project.md"
            selected_project_path.write_text(
                "---\ntitle: Selected Work\n---\n\n## Project | 2026\n\nDescription.\n\n- Result.\n\nTech: Python\n",
                encoding="utf-8",
            )

            markdown = build_markdown_from_source(source)

        self.assertIn("# Davi Naizer", markdown)
        self.assertIn("**Senior Frontend & Product Engineer**", markdown)
        self.assertIn("# Professional Profile", markdown)
        self.assertIn("# Selected Work", markdown)
        self.assertLess(
            markdown.index("# Professional Profile"), markdown.index("# Core Skills")
        )
        self.assertLess(
            markdown.index("# Core Skills"), markdown.index("# Professional Experience")
        )
        self.assertLess(
            markdown.index("# Professional Experience"),
            markdown.index("# Selected Work"),
        )
        self.assertLess(
            markdown.index("# Selected Work"), markdown.index("# Education")
        )

    def test_editable_title_is_passed_to_model_and_renderer_without_changing_section_type(
        self,
    ) -> None:
        titles = read_section_titles(RESUME_SOURCE)
        with tempfile.TemporaryDirectory() as temp_dir:
            source = copy_resume_source(temp_dir)
            summary_path = source / "sections" / "summary.md"
            summary_markdown = summary_path.read_text(encoding="utf-8").replace(
                f"title: {titles['summary']}", "title: Professional Profile", 1
            )
            summary_path.write_text(summary_markdown, encoding="utf-8")

            content = parse_resume_source(source)
            document = build_doc_from_source(source)

        self.assertEqual(content.section_titles.summary, "Professional Profile")
        self.assertEqual(
            heading_texts(document),
            [
                "PROFESSIONAL PROFILE",
                titles["core_skills"].upper(),
                titles["professional_experience"].upper(),
                titles["education"].upper(),
            ],
        )
